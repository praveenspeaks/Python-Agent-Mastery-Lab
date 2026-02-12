import os
from fpdf import FPDF

def create_test_files(directory):
    if not os.path.exists(directory):
        os.makedirs(directory)
    
    # 1. Create a Text File
    with open(os.path.join(directory, "sample_text.txt"), "w") as f:
        f.write("Artificial Intelligence in 2024\n")
        f.write("AI continues to evolve at a breakneck pace. Large Language Models (LLMs) are now capable of multi-modal reasoning.\n")
        f.write("This text file serves as a test case for basic text parsing in our LangChain project.\n")
        f.write("It contains multiple sentences to see how the chunking mechanism splits them based on character limits and overlap.\n")

    # 2. Create a Markdown File
    with open(os.path.join(directory, "sample_markdown.md"), "w") as f:
        f.write("# LangChain Project Documentation\n\n")
        f.write("## Introduction\n")
        f.write("LangChain is a framework for developing applications powered by language models.\n\n")
        f.write("## Core Components\n")
        f.write("- **Loaders**: Bring in data from various sources.\n")
        f.write("- **Splitters**: Break down large documents into chunks.\n")
        f.write("- **Chains**: Combine multiple components to solve a task.\n")

    # 3. Create a PDF File (same as before)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12) # Use standard font
    pdf.cell(200, 10, txt="LangChain PDF Test Document", ln=True, align='C')
    pdf.ln(10)
    text = "This is a sample PDF document for testing LangChain loaders."
    pdf.multi_cell(0, 10, txt=text)
    pdf.output(os.path.join(directory, "sample_pdf.pdf"))

    # 4. Create a CSV File
    import pandas as pd
    df = pd.DataFrame({
        'name': ['Alice', 'Bob', 'Charlie'],
        'age': [25, 30, 35],
        'city': ['London', 'New York', 'Tokyo']
    })
    df.to_csv(os.path.join(directory, "sample_csv.csv"), index=False)

    # 5. Create an Excel File
    df.to_excel(os.path.join(directory, "sample_excel.xlsx"), index=False)

    # 6. Create a Word File
    try:
        from docx import Document as WordDoc
        doc = WordDoc()
        doc.add_heading('LangChain Word Test', 0)
        doc.add_paragraph('This is a test Word document for our parser.')
        doc.save(os.path.join(directory, "sample_word.docx"))
    except ImportError:
        print("python-docx not installed, skipping Word creation.")

    # 7. Create a JSON File
    import json
    data = [
        {"id": 1, "text": "First record for JSON testing."},
        {"id": 2, "text": "Second record for JSON testing."}
    ]
    with open(os.path.join(directory, "sample_json.json"), "w") as f:
        json.dump(data, f)

    # 8. Create an HTML File
    html_content = """
    <html>
        <head><title>Test Page</title></head>
        <body>
            <h1>LangChain HTML Parsing</h1>
            <p>This is a sample paragraph within an HTML document used for testing the BSHTMLLoader.</p>
            <ul>
                <li>Point 1: HTML structure is preserved.</li>
                <li>Point 2: Classes are modular.</li>
            </ul>
        </body>
    </html>
    """
    with open(os.path.join(directory, "sample_html.html"), "w") as f:
        f.write(html_content)

    # 9. Create an XML File
    xml_content = """<?xml version="1.0" encoding="UTF-8"?>
    <note>
        <to>User</to>
        <from>Antigravity</from>
        <heading>Reminder</heading>
        <body>LangChain parsing is powerful and supports XML formats via Unstructured loaders!</body>
    </note>
    """
    with open(os.path.join(directory, "sample_xml.xml"), "w") as f:
        f.write(xml_content)

    print(f"Test data created in: {directory}")

if __name__ == "__main__":
    create_test_files("testdata")
